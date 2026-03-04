"""
ファイル自動仕分けツール
========================
Ollama（gemma3:4b）を使い、ドラッグ＆ドロップしたファイルを
プロジェクトフォルダへ自動分類するデスクトップアプリ。

依存ライブラリ:
    pip install PySide6 pdfplumber python-docx openpyxl

動作環境:
    - Windows 10/11 (64bit)
    - Python 3.13+
    - Ollama（http://localhost:11434 でgemma3:4bが起動済みであること）
"""

import json
import os
import re
import shutil
import sys
import urllib.error
import urllib.request
from datetime import datetime
from pathlib import Path

from PySide6.QtCore import (
    QMimeData,
    QObject,
    QSettings,
    Qt,
    QThread,
    QTimer,
    QUrl,
    Signal,
)
from PySide6.QtGui import QDragEnterEvent, QDropEvent, QFont, QIcon
from PySide6.QtWidgets import (
    QApplication,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

# ---------------------------------------------------------------------------
# 定数
# ---------------------------------------------------------------------------
APP_NAME = "ファイル自動仕分けツール"
APP_VERSION = "1.1.0"
OLLAMA_URL = "http://localhost:11434/api/generate"
OLLAMA_MODEL = "gemma3:4b"
OLLAMA_TIMEOUT = 60
EXTRACT_CHAR_LIMIT = 600
UNKNOWN_DIR = "_未分類"
SUBFOLDERS = ("検討資料", "受領資料", "その他")

# カラーテーマ
C_BG = "#0f0f13"
C_PANEL = "#17171f"
C_BORDER = "#2a2a38"
C_ACCENT_BLUE = "#5b8fff"
C_ACCENT_PURPLE = "#a78bfa"
C_SUCCESS = "#34d399"
C_WARN = "#fbbf24"
C_ERROR = "#f87171"
C_TEXT = "#e8e8f0"
C_SUBTEXT = "#7070a0"

# テキスト抽出可能な拡張子
TEXT_EXTENSIONS = {".txt", ".csv", ".log", ".md", ".json", ".xml", ".html", ".htm"}

# ---------------------------------------------------------------------------
# テキスト抽出
# ---------------------------------------------------------------------------

def extract_text(filepath: str) -> str:
    """ファイルからテキストを抽出する。ライブラリ未導入時はスキップ。"""
    ext = Path(filepath).suffix.lower()

    if ext == ".pdf":
        return _extract_pdf(filepath)
    elif ext == ".docx":
        return _extract_docx(filepath)
    elif ext in (".xlsx", ".xls"):
        return _extract_xlsx(filepath)
    elif ext in TEXT_EXTENSIONS:
        return _extract_plain(filepath)
    else:
        return ""


def _extract_pdf(filepath: str) -> str:
    try:
        import pdfplumber
    except ImportError:
        return ""
    try:
        text_parts: list[str] = []
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages[:3]:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
        return "\n".join(text_parts)[:EXTRACT_CHAR_LIMIT]
    except Exception:
        return ""


def _extract_docx(filepath: str) -> str:
    try:
        import docx
    except ImportError:
        return ""
    try:
        doc = docx.Document(filepath)
        text_parts = [p.text for p in doc.paragraphs[:20] if p.text.strip()]
        return "\n".join(text_parts)[:EXTRACT_CHAR_LIMIT]
    except Exception:
        return ""


def _extract_xlsx(filepath: str) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        rows: list[str] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 11:
                break
            cells = [str(c) if c is not None else "" for c in row]
            rows.append("\t".join(cells))
        wb.close()
        return "\n".join(rows)[:EXTRACT_CHAR_LIMIT]
    except Exception:
        return ""


def _extract_plain(filepath: str) -> str:
    encodings = ("utf-8", "cp932", "shift_jis", "euc-jp", "latin-1")
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                return f.read(EXTRACT_CHAR_LIMIT)
        except (UnicodeDecodeError, LookupError):
            continue
    return ""


# ---------------------------------------------------------------------------
# プロジェクトフォルダ検出
# ---------------------------------------------------------------------------

def scan_projects(root: str) -> list[str]:
    """ルートフォルダ直下のディレクトリ名を取得。先頭が _ のものは除外。"""
    root_path = Path(root)
    if not root_path.is_dir():
        return []
    return sorted(
        d.name
        for d in root_path.iterdir()
        if d.is_dir() and not d.name.startswith("_") and not d.name.startswith(".")
    )


# ---------------------------------------------------------------------------
# Ollama 呼び出し
# ---------------------------------------------------------------------------

def build_prompt(filename: str, content: str, projects: list[str]) -> str:
    content_display = content.strip() if content.strip() else "（読み取り不可）"
    project_list = "\n".join(f"- {p}" for p in projects) if projects else "（なし）"

    return f"""あなたはファイル整理の専門家です。以下の情報をもとに、ファイルの分類先を答えてください。

## ファイル名
{filename}

## ファイルの内容（先頭抜粋）
{content_display}

## 既存プロジェクトフォルダ一覧
{project_list}

## 分類ルール
- projectには既存プロジェクトフォルダ一覧のいずれか、または "unknown" を指定
- サブフォルダは必ず「検討資料」「受領資料」「その他」のいずれか1つ
  - 検討資料：自分たちが作成・編集する設計図・計算書・報告書・提案書など
  - 受領資料：発注者・他社・官庁から受け取った資料・データ・提供ファイルなど
  - その他：議事録・写真・メモ・分類が難しいもの

## 回答形式（JSONのみ・余計な文字禁止）
{{"project": "プロジェクト名またはunknown", "subfolder": "検討資料 or 受領資料 or その他"}}"""


def call_ollama(prompt: str) -> dict | None:
    """Ollama API を呼び出し、JSON をパースして返す。"""
    payload = json.dumps({
        "model": OLLAMA_MODEL,
        "prompt": prompt,
        "stream": False,
        "options": {"temperature": 0.1},
    }).encode("utf-8")

    req = urllib.request.Request(
        OLLAMA_URL,
        data=payload,
        headers={"Content-Type": "application/json"},
    )

    try:
        with urllib.request.urlopen(req, timeout=OLLAMA_TIMEOUT) as resp:
            body = json.loads(resp.read().decode("utf-8"))
    except (urllib.error.URLError, TimeoutError, json.JSONDecodeError):
        return None

    raw = body.get("response", "")
    # JSON 部分を抽出（```json ... ``` やテキスト混入に対応）
    json_match = re.search(r"\{[^}]*\"project\"[^}]*\}", raw, re.DOTALL)
    if not json_match:
        return None
    try:
        return json.loads(json_match.group())
    except json.JSONDecodeError:
        return None


def check_ollama_connection() -> bool:
    """Ollama サーバーへの接続を確認する。"""
    try:
        req = urllib.request.Request("http://localhost:11434/api/tags")
        with urllib.request.urlopen(req, timeout=5):
            return True
    except Exception:
        return False


# ---------------------------------------------------------------------------
# ファイル移動
# ---------------------------------------------------------------------------

def move_file(filepath: str, dest_dir: str) -> str:
    """ファイルを移動する。同名ファイルがあればタイムスタンプを付与。
    移動後のパスを返す。"""
    dest_path = Path(dest_dir)
    dest_path.mkdir(parents=True, exist_ok=True)

    src = Path(filepath)
    target = dest_path / src.name

    if target.exists():
        stem = src.stem
        suffix = src.suffix
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        target = dest_path / f"{stem}_{ts}{suffix}"

    shutil.move(str(src), str(target))
    return str(target)


# ---------------------------------------------------------------------------
# ワーカースレッド
# ---------------------------------------------------------------------------

class SortWorker(QObject):
    """バックグラウンドでファイル仕分けを実行するワーカー。"""

    log_signal = Signal(str, str)  # (message, color)
    status_signal = Signal(str)
    finished = Signal()
    undo_record = Signal(str, str)  # (moved_to, original_path)

    def __init__(self, files: list[str], root_folder: str):
        super().__init__()
        self.files = files
        self.root_folder = root_folder

    def run(self):
        projects = scan_projects(self.root_folder)

        for filepath in self.files:
            filename = Path(filepath).name
            self.status_signal.emit(f"処理中… {filename}")
            self.log_signal.emit(f"🔍 解析中: {filename}", C_SUBTEXT)

            # テキスト抽出
            content = extract_text(filepath)

            # Ollama 呼び出し
            prompt = build_prompt(filename, content, projects)
            result = call_ollama(prompt)

            if result is None:
                # パース失敗 → 未分類
                dest_dir = str(Path(self.root_folder) / UNKNOWN_DIR)
                moved = move_file(filepath, dest_dir)
                self.undo_record.emit(moved, filepath)
                self.log_signal.emit(
                    f"⚠️ AI応答を解析できません → {UNKNOWN_DIR}/ へ移動: {filename}",
                    C_WARN,
                )
                continue

            project = result.get("project", "unknown")
            subfolder = result.get("subfolder", "その他")

            # サブフォルダのバリデーション
            if subfolder not in SUBFOLDERS:
                subfolder = "その他"

            if project == "unknown" or project not in projects:
                dest_dir = str(Path(self.root_folder) / UNKNOWN_DIR)
                moved = move_file(filepath, dest_dir)
                self.undo_record.emit(moved, filepath)
                self.log_signal.emit(
                    f"⚠️ プロジェクト不明 → {UNKNOWN_DIR}/ へ移動: {filename}",
                    C_WARN,
                )
            else:
                dest_dir = str(Path(self.root_folder) / project / subfolder)
                moved = move_file(filepath, dest_dir)
                self.undo_record.emit(moved, filepath)
                self.log_signal.emit(
                    f"✅ {filename} → {project}/{subfolder}/",
                    C_SUCCESS,
                )

        self.status_signal.emit("✅ 完了")
        self.finished.emit()


# ---------------------------------------------------------------------------
# メインウィンドウ
# ---------------------------------------------------------------------------

class DropZone(QFrame):
    """ドラッグ＆ドロップを受け付けるエリア。"""

    files_dropped = Signal(list)

    def __init__(self):
        super().__init__()
        self.setAcceptDrops(True)
        self.setMinimumHeight(140)
        self._update_style(False)

        layout = QVBoxLayout(self)
        layout.setAlignment(Qt.AlignCenter)

        icon_label = QLabel("📂")
        icon_label.setAlignment(Qt.AlignCenter)
        icon_label.setStyleSheet(f"font-size: 36px; border: none; background: transparent;")
        layout.addWidget(icon_label)

        text_label = QLabel("ここにファイルをドロップ（複数可）")
        text_label.setAlignment(Qt.AlignCenter)
        text_label.setStyleSheet(
            f"color: {C_SUBTEXT}; font-size: 14px; border: none; background: transparent;"
        )
        layout.addWidget(text_label)

    def _update_style(self, hover: bool):
        border_color = C_ACCENT_BLUE if hover else C_BORDER
        bg = "#1c1c2a" if hover else C_PANEL
        self.setStyleSheet(
            f"DropZone {{ background: {bg}; border: 2px dashed {border_color}; "
            f"border-radius: 12px; }}"
        )

    def dragEnterEvent(self, event: QDragEnterEvent):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
            self._update_style(True)

    def dragLeaveEvent(self, event):
        self._update_style(False)

    def dropEvent(self, event: QDropEvent):
        self._update_style(False)
        urls: list[QUrl] = event.mimeData().urls()
        files = [u.toLocalFile() for u in urls if Path(u.toLocalFile()).is_file()]
        if files:
            self.files_dropped.emit(files)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"{APP_NAME} v{APP_VERSION}")
        self.resize(680, 720)
        self.setMinimumSize(520, 580)

        self._settings = QSettings("FileSorter", "FileSorter")
        self._worker_thread: QThread | None = None
        self._undo_stack: list[tuple[str, str]] = []  # (moved_to, original_path)

        self._init_ui()
        self._apply_global_style()
        self._load_settings()
        self._refresh_projects()
        self._check_ollama()

    # ---- UI 構築 ----

    def _init_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root_layout = QVBoxLayout(central)
        root_layout.setContentsMargins(20, 16, 20, 16)
        root_layout.setSpacing(12)

        # ヘッダー
        header = QLabel(APP_NAME)
        header.setStyleSheet(
            f"color: {C_TEXT}; font-size: 22px; font-weight: bold;"
        )
        header.setAlignment(Qt.AlignCenter)
        root_layout.addWidget(header)

        sub = QLabel(f"Powered by Ollama + {OLLAMA_MODEL}・完全ローカル処理")
        sub.setStyleSheet(f"color: {C_SUBTEXT}; font-size: 12px;")
        sub.setAlignment(Qt.AlignCenter)
        root_layout.addWidget(sub)

        # Ollama 接続状態
        self._ollama_status = QLabel()
        self._ollama_status.setAlignment(Qt.AlignCenter)
        self._ollama_status.setStyleSheet(f"font-size: 11px;")
        root_layout.addWidget(self._ollama_status)

        # フォルダ選択バー
        folder_bar = QHBoxLayout()
        folder_label = QLabel("整理先フォルダ:")
        folder_label.setStyleSheet(f"color: {C_SUBTEXT}; font-size: 12px;")
        folder_bar.addWidget(folder_label)

        self._folder_edit = QLineEdit()
        self._folder_edit.setReadOnly(True)
        self._folder_edit.setStyleSheet(
            f"background: {C_PANEL}; color: {C_TEXT}; border: 1px solid {C_BORDER}; "
            f"border-radius: 6px; padding: 4px 8px; font-size: 12px;"
        )
        folder_bar.addWidget(self._folder_edit, 1)

        change_btn = QPushButton("変更")
        change_btn.setFixedWidth(64)
        change_btn.setCursor(Qt.PointingHandCursor)
        change_btn.setStyleSheet(self._button_style(C_ACCENT_PURPLE))
        change_btn.clicked.connect(self._select_folder)
        folder_bar.addWidget(change_btn)

        root_layout.addLayout(folder_bar)

        # プロジェクト一覧
        self._projects_label = QLabel()
        self._projects_label.setWordWrap(True)
        self._projects_label.setStyleSheet(f"color: {C_SUBTEXT}; font-size: 11px;")
        root_layout.addWidget(self._projects_label)

        # ドロップゾーン
        self._drop_zone = DropZone()
        self._drop_zone.files_dropped.connect(self._on_files_dropped)
        root_layout.addWidget(self._drop_zone)

        # ステータス
        self._status_label = QLabel("待機中")
        self._status_label.setAlignment(Qt.AlignCenter)
        self._status_label.setStyleSheet(f"color: {C_ACCENT_BLUE}; font-size: 13px;")
        root_layout.addWidget(self._status_label)

        # ログエリア
        self._log_area = QTextEdit()
        self._log_area.setReadOnly(True)
        self._log_area.setMinimumHeight(160)
        self._log_area.setStyleSheet(
            f"background: {C_PANEL}; color: {C_TEXT}; border: 1px solid {C_BORDER}; "
            f"border-radius: 8px; padding: 8px; font-size: 12px;"
        )
        root_layout.addWidget(self._log_area, 1)

        # ボタンバー
        btn_bar = QHBoxLayout()

        undo_btn = QPushButton("↩ 元に戻す")
        undo_btn.setCursor(Qt.PointingHandCursor)
        undo_btn.setStyleSheet(self._button_style(C_WARN))
        undo_btn.clicked.connect(self._undo_last)
        btn_bar.addWidget(undo_btn)

        btn_bar.addStretch()

        clear_btn = QPushButton("ログをクリア")
        clear_btn.setCursor(Qt.PointingHandCursor)
        clear_btn.setStyleSheet(self._button_style(C_ACCENT_BLUE))
        clear_btn.clicked.connect(self._log_area.clear)
        btn_bar.addWidget(clear_btn)

        root_layout.addLayout(btn_bar)

    # ---- スタイル ----

    def _apply_global_style(self):
        self.setStyleSheet(
            f"QMainWindow {{ background: {C_BG}; }}"
            f"QWidget {{ background: {C_BG}; color: {C_TEXT}; font-family: 'Yu Gothic UI', 'Meiryo', 'Segoe UI', sans-serif; }}"
            f"QScrollBar:vertical {{ background: {C_PANEL}; width: 8px; border-radius: 4px; }}"
            f"QScrollBar::handle:vertical {{ background: {C_BORDER}; border-radius: 4px; min-height: 30px; }}"
            f"QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{ height: 0px; }}"
        )

    @staticmethod
    def _button_style(color: str) -> str:
        return (
            f"QPushButton {{ background: {color}20; color: {color}; "
            f"border: 1px solid {color}60; border-radius: 6px; "
            f"padding: 6px 14px; font-size: 12px; font-weight: bold; }}"
            f"QPushButton:hover {{ background: {color}40; }}"
            f"QPushButton:pressed {{ background: {color}60; }}"
        )

    # ---- 設定 ----

    def _load_settings(self):
        default = str(Path.home() / "Desktop")
        folder = self._settings.value("root_folder", default)
        if not Path(folder).is_dir():
            folder = default
        self._folder_edit.setText(folder)

    def _save_settings(self):
        self._settings.setValue("root_folder", self._folder_edit.text())

    @property
    def root_folder(self) -> str:
        return self._folder_edit.text()

    # ---- Ollama 接続確認 ----

    def _check_ollama(self):
        connected = check_ollama_connection()
        if connected:
            self._ollama_status.setText("🟢 Ollama 接続済み")
            self._ollama_status.setStyleSheet(f"color: {C_SUCCESS}; font-size: 11px;")
        else:
            self._ollama_status.setText("🔴 Ollama 未接続 — サーバーを起動してください")
            self._ollama_status.setStyleSheet(f"color: {C_ERROR}; font-size: 11px;")

    # ---- フォルダ選択 ----

    def _select_folder(self):
        folder = QFileDialog.getExistingDirectory(
            self, "整理先ルートフォルダを選択", self.root_folder
        )
        if folder:
            self._folder_edit.setText(folder)
            self._save_settings()
            self._refresh_projects()

    def _refresh_projects(self):
        projects = scan_projects(self.root_folder)
        if projects:
            self._projects_label.setText(
                f"検出プロジェクト ({len(projects)}): " + "、".join(projects)
            )
        else:
            self._projects_label.setText("プロジェクトフォルダが見つかりません")

    # ---- ファイルドロップ処理 ----

    def _on_files_dropped(self, files: list[str]):
        if self._worker_thread is not None and self._worker_thread.isRunning():
            QMessageBox.warning(self, "処理中", "処理中です。完了をお待ちください。")
            return

        # Ollama 再確認
        if not check_ollama_connection():
            self._check_ollama()
            self._append_log("❌ Ollama に接続できません。サーバーを起動してください。", C_ERROR)
            return

        self._refresh_projects()
        self._status_label.setText(f"処理中… (0/{len(files)})")

        # ワーカー起動
        self._worker_thread = QThread()
        self._worker = SortWorker(files, self.root_folder)
        self._worker.moveToThread(self._worker_thread)

        self._worker.log_signal.connect(self._append_log)
        self._worker.status_signal.connect(self._status_label.setText)
        self._worker.finished.connect(self._on_worker_finished)
        self._worker.undo_record.connect(self._record_undo)

        self._worker_thread.started.connect(self._worker.run)
        self._worker_thread.start()

    def _on_worker_finished(self):
        if self._worker_thread:
            self._worker_thread.quit()
            self._worker_thread.wait()
            self._worker_thread = None
        self._check_ollama()

    # ---- 元に戻す ----

    def _record_undo(self, moved_to: str, original_path: str):
        self._undo_stack.append((moved_to, original_path))

    def _undo_last(self):
        if not self._undo_stack:
            self._append_log("↩ 戻せる操作がありません。", C_SUBTEXT)
            return

        moved_to, original_path = self._undo_stack.pop()
        moved_path = Path(moved_to)

        if not moved_path.exists():
            self._append_log(f"❌ ファイルが見つかりません: {moved_to}", C_ERROR)
            return

        try:
            original_dir = str(Path(original_path).parent)
            Path(original_dir).mkdir(parents=True, exist_ok=True)
            shutil.move(str(moved_path), original_path)
            self._append_log(
                f"↩ 元に戻しました: {moved_path.name} → {original_path}",
                C_ACCENT_PURPLE,
            )
        except Exception as e:
            self._append_log(f"❌ 元に戻す操作に失敗: {e}", C_ERROR)

    # ---- ログ ----

    def _append_log(self, message: str, color: str = C_TEXT):
        timestamp = datetime.now().strftime("%H:%M:%S")
        self._log_area.append(
            f'<span style="color:{C_SUBTEXT}">[{timestamp}]</span> '
            f'<span style="color:{color}">{message}</span>'
        )
        # 最下部へスクロール
        sb = self._log_area.verticalScrollBar()
        sb.setValue(sb.maximum())


# ---------------------------------------------------------------------------
# エントリーポイント
# ---------------------------------------------------------------------------

def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")

    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
